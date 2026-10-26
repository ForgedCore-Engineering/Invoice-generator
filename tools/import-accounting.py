"""Read-only XLSX extraction. Emits an additive, repeat-safe SQL seed and review manifest.
Usage: python tools/import-accounting.py workbook.xlsx
Requires openpyxl (reading only). Never executes workbook content or modifies the source.
"""
import sys, json, hashlib, re, datetime, pathlib, shutil, collections, ast
from decimal import Decimal
import openpyxl

root = pathlib.Path(__file__).resolve().parents[1]
source = pathlib.Path(sys.argv[1])
formulas = openpyxl.load_workbook(source, data_only=False)
values = openpyxl.load_workbook(source, data_only=True)
digest = hashlib.sha256(source.read_bytes()).hexdigest()
out = root / 'accounting-private'
out.mkdir(exist_ok=True)
shutil.copy2(source, out / 'original.xlsx')
tables = collections.defaultdict(list)
issues = []
def add(table, **row):
    tables[table].append(row)
def v(sheet, addr):
    return values[sheet][addr].value
def formula_value(sheet, formula):
    # Strict arithmetic parser for this workbook's SUM/PRODUCT formulas, never eval().
    # Resolve reference ranges before parsing, including the original full-column SUM.
    def cell_value(addr):
        value=v(sheet,addr)
        return Decimal(str(value)) if isinstance(value,(int,float,Decimal)) else Decimal(0)
    def range_sum(match):
        a,b=match.group(1).replace('$',''),match.group(2).replace('$','')
        if a.isalpha():a+='1'
        if b.isalpha():b+=str(values[sheet].max_row)
        cells=values[sheet][a:b]
        return str(sum((Decimal(str(c.value)) for row in cells for c in row if isinstance(c.value,(int,float,Decimal))),Decimal(0)))
    expression=re.sub(r'(\$?[A-Z]+\$?\d*):(\$?[A-Z]+\$?\d*)',range_sum,formula.lstrip('='))
    expression=re.sub(r'\$?[A-Z]+\$?\d+',lambda m:str(cell_value(m.group().replace('$',''))),expression)
    def calc(node):
        if isinstance(node,ast.Expression):return calc(node.body)
        if isinstance(node,ast.Constant) and isinstance(node.value,(int,float)):return Decimal(str(node.value))
        if isinstance(node,ast.UnaryOp) and isinstance(node.op,ast.USub):return -calc(node.operand)
        if isinstance(node,ast.BinOp):
            a,b=calc(node.left),calc(node.right)
            if isinstance(node.op,ast.Add):return a+b
            if isinstance(node.op,ast.Sub):return a-b
            if isinstance(node.op,ast.Mult):return a*b
            if isinstance(node.op,ast.Div):return a/b
        if isinstance(node,ast.Call) and isinstance(node.func,ast.Name) and not node.keywords:
            args=[calc(x) for x in node.args]
            if node.func.id=='SUM':return sum(args,Decimal(0))
            if node.func.id=='PRODUCT':
                result=Decimal(1)
                for x in args:result*=x
                return result
        raise ValueError('Unsupported formula expression: '+formula)
    return calc(ast.parse(expression,mode='eval'))
def text(x):
    return '' if x is None else str(x)
def date(x, ref):
    if x is None: return None
    if isinstance(x, (datetime.datetime, datetime.date)): return x.strftime('%Y-%m-%d')
    s = str(x).strip().lower()
    for a,b in {'augist':'august','agust':'august','octber':'october','octorber':'october','octoober':'october'}.items(): s=s.replace(a,b)
    s=re.sub(r'(\d+)(st|nd|rd|th)',r'\1',s)
    for fmt in ('%d %B %Y','%d/%m/%Y','%Y-%m-%d'):
        try:return datetime.datetime.strptime(s,fmt).strftime('%Y-%m-%d')
        except ValueError:pass
    issues.append({'reference':ref,'issue':'Unresolved date preserved as text','value':text(x)})
    return None

s='Project Tracker'
for row in list(range(3,8))+list(range(10,33)):
    if not v(s,f'A{row}'):continue
    add('projects',name=v(s,f'A{row}'),kind='Domain' if row<8 else 'Project',
        start_date=date(v(s,f'B{row}'),f'{s}!B{row}'),expiry_date=date(v(s,f'C{row}'),f'{s}!C{row}'),
        handler=text(v(s,f'D{row}')),department=text(v(s,f'E{row}')),priority=text(v(s,f'F{row}')),
        status=text(v(s,f'G{row}')),notes=text(v(s,f'H{row}')),source_ref=f'{s}!A{row}')
s='Financial Sheet'
legacy_expected=set(re.findall(r'B(\d+)',formulas[s]['B158'].value))
legacy_received=set(re.findall(r'C(\d+)',formulas[s]['D159'].value))
current=None
for row in range(4,157):
    name=v(s,f'A{row}'); expected=v(s,f'B{row}'); paid=v(s,f'C{row}')
    if name in ('Outstanding','New Year Deals'): name=None
    # A blank name with a new expected value is a distinct work item, not an added fee.
    if expected is not None or (name and name != 'House Renovation construction'):
        label=text(name) or (tables['engagements'][current-1]['name'] if current else 'Unassigned')
        if not name:label += ' — '+text(v(s,f'D{row}'))+' (expected amount entry)'
        current=len(tables['engagements'])+1
        note=text(v(s,f'D{row}'))
        if row==84:
            note+='; B84 may revise B83 or represent separate work. Kept separate for review; no amounts discarded.'
            issues.append({'reference':f'{s}!B83:B84','issue':'Two expected amounts for lighting project; review whether B84 revises B83'})
        add('engagements',id=current,name=label,expected=expected,legacy_expected=int(str(row) in legacy_expected),notes=note,source_ref=f'{s}!A{row}:E{row}')
    if isinstance(paid,(int,float,Decimal)) and row!=16:
        if current is None: raise ValueError('Payment without engagement')
        add('payments',engagement_id=current,amount=paid,payment_date=date(v(s,f'E{row}'),f'{s}!E{row}'),
            date_original=text(v(s,f'E{row}')),notes=text(v(s,f'D{row}')),legacy_received=int(str(row) in legacy_received),source_ref=f'{s}!C{row}')
    spent=v(s,f'F{row}')
    if isinstance(spent,(int,float,Decimal)):
        add('expenses',amount=spent,expense_date=date(v(s,f'I{row}'),f'{s}!I{row}'),date_original=text(v(s,f'I{row}')),
            description=text(v(s,f'G{row}')),payment_details=text(v(s,f'H{row}')),notes='',source_ref=f'{s}!F{row}')
s='Income Breakdown'
for header,base,status,dept in [(3,4,10,'Software Engineering'),(17,18,23,'Architecture, Construction & MEP')]:
    for col in range(3,27):
        letter=openpyxl.utils.get_column_letter(col);name=v(s,f'{letter}{header}')
        if not name:continue
        amount=v(s,f'{letter}{base}')
        data=dict(name=name,department=dept,amount=amount,status=text(v(s,f'{letter}{status}')),notes='',source_ref=f'{s}!{letter}{base}')
        for offset,key in enumerate(['sourcer','dividend','department','holdings'],1):
            rate=v(s,f'B{base+offset}');value=v(s,f'{letter}{base+offset}')
            data[key+'_rate']=rate
            data[key+'_override']=value if formulas[s][f'{letter}{base+offset}'].data_type!='f' else None
            if data[key+'_override'] is not None:
                data['notes']+='Original fixed '+key+' allocation retained. '
        add('allocations',**data)
add('reserves',name='USDT savings (workbook cash deduction)',amount=v('Financial Sheet','I161'),notes='Preserved as the workbook cash deduction. No currency conversion inferred.',source_ref='Financial Sheet!I161')
for row in range(2,13):
    add('assets',name=v('Company Assets Register',f'A{row}'),quantity=v('Company Assets Register',f'B{row}'),specification=text(v('Company Assets Register',f'C{row}')),source_ref=f'Company Assets Register!A{row}')
for row in range(3,13):
    number=v('SIM cards and Serials',f'A{row}')
    add('sims',sim_number=format(number,'.0f') if isinstance(number,(int,float)) else text(number),service_id=text(v('SIM cards and Serials',f'B{row}')),notes='Number retained without inventing a leading zero.',source_ref=f'SIM cards and Serials!A{row}')
cells=[]
formula_checks=0
for ws in formulas:
    for row in ws:
        for c in row:
            if c.value is None:continue
            value=values[ws.title][c.coordinate].value
            cells.append(dict(sheet=ws.title,address=c.coordinate,value_text=json.dumps(value,default=str,ensure_ascii=False),formula=c.value if c.data_type=='f' else None,cell_type=c.data_type))
            if c.data_type=='f':
                computed=formula_value(ws.title,c.value)
                assert abs(computed-Decimal(str(value)))<Decimal('.00001'),f'Formula mismatch: {ws.title}!{c.coordinate}'
                formula_checks+=1
totals={
 'expected_all':sum(Decimal(str(r['expected'] or 0)) for r in tables['engagements']),
 'received_all':sum(Decimal(str(r['amount'])) for r in tables['payments']),
 'expected_legacy':sum(Decimal(str(r['expected'] or 0)) for r in tables['engagements'] if r['legacy_expected']),
 'received_legacy':sum(Decimal(str(r['amount'])) for r in tables['payments'] if r['legacy_received']),
 'expenses':sum(Decimal(str(r['amount'])) for r in tables['expenses'])}
assert totals['expected_legacy']==Decimal(str(v('Financial Sheet','B158')))
assert totals['received_legacy']==Decimal(str(v('Financial Sheet','D159')))
assert totals['expenses']==Decimal(str(v('Financial Sheet','G159')))
issues += [
 {'reference':'Financial Sheet!B158,D159','issue':'Original totals include selected entries only. All-entry reports include excluded projects and construction payments. Gross construction receipts and profit deposits may overlap; preserve original reporting scope until reviewed.'},
 {'reference':'Income Breakdown!J19','issue':'Fixed 22000 sourcer allocation exceeds 5%; preserved and allocation variance exposed.'},
 {'reference':'Income Breakdown!J29, Financial Sheet!J158','issue':'Unlabelled amounts preserved in original cell archive; not classified as transactions.'}]
manifest=dict(filename=source.name,sha256=digest,sheets=formulas.sheetnames,counts={k:len(x) for k,x in tables.items()},source_cells=len(cells),formula_count=sum(c['formula'] is not None for c in cells),verified_formulas=formula_checks,totals=totals,issues=issues)
(out/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
def q(x):
    if x is None:return 'NULL'
    if isinstance(x,(int,float,Decimal)):return str(x)
    # Hex strings avoid SQL mode-dependent escaping of quotes/backslashes.
    return "CONVERT(X'"+str(x).encode('utf-8').hex()+"' USING utf8mb4)"
sql=['-- Generated from the preserved workbook; additive and safe to run again.', 'START TRANSACTION;',
     'INSERT IGNORE INTO acct_imports (source_hash,filename,manifest) VALUES ('+','.join(map(q,[digest,source.name,json.dumps(manifest,default=str,ensure_ascii=False)]))+');',
     'SET @acct_import_id=(SELECT id FROM acct_imports WHERE source_hash='+q(digest)+');']
for c in cells:
    sql.append('INSERT IGNORE INTO acct_source_cells (import_id,'+','.join(c)+') VALUES (@acct_import_id,'+','.join(q(x) for x in c.values())+');')
for table,rows in tables.items():
    for data in rows:
        data=dict(data)
        # Resolve parent by stable source reference rather than numeric IDs in the target DB.
        data.pop('id',None)
        parent=None
        if table=='payments':
            parent=tables['engagements'][data.pop('engagement_id')-1]['source_ref']
        columns=list(data)
        expressions=[q(x) for x in data.values()]
        if parent:
            columns.insert(0,'engagement_id');expressions.insert(0,'(SELECT id FROM acct_engagements WHERE source_ref='+q(parent)+')')
        sql.append('INSERT INTO acct_'+table+' ('+','.join(columns)+') SELECT '+','.join(expressions)+' WHERE NOT EXISTS (SELECT 1 FROM acct_'+table+' WHERE source_ref='+q(data['source_ref'])+');')
sql.append('COMMIT;')
(out/'accounting_data.sql').write_text('\n'.join(sql)+'\n',encoding='utf-8')
(out/'accounting_install.sql').write_text((root/'accounting_schema.sql').read_text(encoding='utf-8')+'\n'+'\n'.join(sql)+'\n',encoding='utf-8')
print(json.dumps(manifest,indent=2,default=str,ensure_ascii=False))
